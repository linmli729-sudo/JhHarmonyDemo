# -*- coding: utf-8 -*-
"""產生《軟體測試》大作業報告封面（書名頁）Word 檔。"""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn


def set_cell_font(cell, size=14, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT):
    for p in cell.paragraphs:
        p.alignment = align
        for r in p.runs:
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.name = "標楷體"
            r._element.rPr.rFonts.set(qn("w:eastAsia"), "標楷體")


def add_title(doc, text, size, bold=True, space_before=0, space_after=12):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.name = "標楷體"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "標楷體")
    return p


doc = Document()
sec = doc.sections[0]
sec.top_margin = Cm(3)
sec.bottom_margin = Cm(2.5)
sec.left_margin = Cm(3)
sec.right_margin = Cm(3)

# 校名
add_title(doc, "台灣國立大學", 26, bold=True, space_before=24, space_after=40)
# 報告標題
add_title(doc, "《軟體測試》", 30, bold=True, space_after=6)
add_title(doc, "大作業報告", 30, bold=True, space_after=30)
# 副標
add_title(doc, "—— 以 Automation Exercise 電子商務網站為例 ——", 14, bold=False, space_after=48)

# 資訊表格
fields = [
    ("課程名稱", "軟體測試"),
    ("被測系統", "Automation Exercise（https://automationexercise.com）"),
    ("姓　　名", "〔手寫〕"),
    ("學　　號", "〔手寫〕"),
    ("班　　級", "〔手寫〕"),
    ("指導老師", "張　導師"),
    ("日　　期", "中華民國　　年　　月　　日"),
]
table = doc.add_table(rows=len(fields), cols=2)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.style = "Table Grid"
table.columns[0].width = Cm(3.5)
table.columns[1].width = Cm(10.5)
for i, (k, v) in enumerate(fields):
    c0 = table.rows[i].cells[0]
    c1 = table.rows[i].cells[1]
    c0.width = Cm(3.5)
    c1.width = Cm(10.5)
    c0.text = k
    c1.text = v
    set_cell_font(c0, size=14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_font(c1, size=14, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT)
    table.rows[i].height = Cm(1.0)

# 備註
note = doc.add_paragraph()
note.alignment = WD_ALIGN_PARAGRAPH.LEFT
note.paragraph_format.space_before = Pt(36)
r = note.add_run("〔封面依老師要求：列印後手寫填寫姓名、學號、班級與日期；指導老師欄已填張導師〕")
r.font.size = Pt(10.5)
r.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
r.font.name = "標楷體"
r._element.rPr.rFonts.set(qn("w:eastAsia"), "標楷體")

out = "軟體測試大作業/封面_書名頁.docx"
doc.save(out)
print("已產生:", out)
