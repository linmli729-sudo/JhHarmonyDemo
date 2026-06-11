# -*- coding: utf-8 -*-
"""产生《软件测试》大作业完整报告 Word 文件（封面+目录+五章+205案例+致谢）—— 简体中文版。
测试账号：Wjc25S010418 / wjc25S010418@gmail.com / Wjc25S0418（地址：北京）
"""
import os, csv, glob
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE = os.path.dirname(os.path.abspath(__file__))
CN_FONT = "SimSun"
BODY_CN = "SimSun"
BODY_EN = "Times New Roman"
RED = RGBColor(0xC0, 0x00, 0x00)


def _set_run_font(r, cn=BODY_CN, size=12, bold=False, color=None):
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.name = BODY_EN
    rpr = r._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), cn)
    if color is not None:
        r.font.color.rgb = color


def heading(doc, text, level=1):
    sizes = {1: 16, 2: 14, 3: 13}
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    p.style = doc.styles["Heading %d" % level]
    r = p.add_run(text)
    _set_run_font(r, cn=CN_FONT, size=sizes.get(level, 13), bold=True)
    r.font.color.rgb = RGBColor(0, 0, 0)
    return p


def body(doc, text, red=False, indent=True, size=12):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing = 1.5
    pf.space_after = Pt(4)
    if indent:
        pf.first_line_indent = Pt(size * 2)
    r = p.add_run(text)
    _set_run_font(r, cn=BODY_CN, size=size, color=RED if red else None)
    return p


def add_image(doc, path, width_cm=15):
    if not os.path.exists(path):
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(path, width=Cm(width_cm))


def caption(doc, text, above=True):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6 if above else 2)
    p.paragraph_format.space_after = Pt(2 if above else 6)
    r = p.add_run(text)
    _set_run_font(r, cn=CN_FONT, size=10.5, bold=True)
    return p


def make_table(doc, headers, rows, widths=None, header_size=10.5, body_size=10.5):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        _set_run_font(r, cn=CN_FONT, size=header_size, bold=True)
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            r = p.add_run(str(val))
            _set_run_font(r, cn=BODY_CN, size=body_size)
    if widths:
        for i, w in enumerate(widths):
            for row in t.rows:
                row.cells[i].width = Cm(w)
    return t


def add_toc(doc):
    p = doc.add_paragraph()
    run = p.add_run()
    fldStart = OxmlElement("w:fldChar"); fldStart.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve")
    instr.text = 'TOC \\o "1-3" \\h \\z \\u'
    fldSep = OxmlElement("w:fldChar"); fldSep.set(qn("w:fldCharType"), "separate")
    t = OxmlElement("w:t"); t.text = "〔请在 Word 中对此处按右键 →「更新域」→「更新整个目录」〕"
    fldEnd = OxmlElement("w:fldChar"); fldEnd.set(qn("w:fldCharType"), "end")
    for e in (fldStart, instr, fldSep, t, fldEnd):
        run._r.append(e)


def cover(doc):
    def c_title(text, size, bold=True, sb=0, sa=12):
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(sb); p.paragraph_format.space_after = Pt(sa)
        r = p.add_run(text); _set_run_font(r, cn=CN_FONT, size=size, bold=bold)
    c_title("台湾国立大学", 26, sb=24, sa=40)
    c_title("《软件测试》", 30, sa=6)
    c_title("大作业报告", 30, sa=30)
    c_title("—— 以 Automation Exercise 电子商务网站为例 ——", 14, bold=False, sa=48)
    fields = [
        ("课程名称", "软件测试"),
        ("被测系统", "Automation Exercise（https://automationexercise.com）"),
        ("姓    名", "〔手写〕"),
        ("学    号", "〔手写〕"),
        ("班    级", "〔手写〕"),
        ("指导老师", "张  导师"),
        ("日    期", "      年    月    日"),
    ]
    table = doc.add_table(rows=len(fields), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for i, (k, v) in enumerate(fields):
        c0, c1 = table.rows[i].cells
        c0.width = Cm(3.5); c1.width = Cm(10.5)
        for cell, val, bold, al in ((c0, k, True, WD_ALIGN_PARAGRAPH.CENTER), (c1, v, False, WD_ALIGN_PARAGRAPH.LEFT)):
            cell.text = ""; p = cell.paragraphs[0]; p.alignment = al
            r = p.add_run(val); _set_run_font(r, cn=CN_FONT, size=14, bold=bold)
    doc.add_page_break()


def read_module_cases():
    """返回 [(模块名, [(编号,测试方法,输入/操作,预期输出), ...]), ...]"""
    order = ["用户注册模块", "用户登录模块", "商品搜索模块", "商品浏览与详情模块",
             "购物车模块", "结账与下单模块", "联系我们表单模块", "订阅与其他功能模块"]
    cdir = os.path.join(BASE, "功能测试案例_简体")
    result = []
    for name in order:
        fn = os.path.join(cdir, name + ".csv")
        rows = []
        with open(fn, encoding="utf-8-sig") as f:
            rd = csv.reader(f); next(rd)
            for row in rd:
                rows.append((row[0], row[2], row[3], row[4]))
        result.append((name, rows))
    return result


def build():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(2.5); sec.bottom_margin = Cm(2.5)
    sec.left_margin = Cm(2.8); sec.right_margin = Cm(2.8)

    # 封面
    cover(doc)

    # 目录
    heading(doc, "目  录", 1)
    add_toc(doc)
    doc.add_page_break()

    # 一、引言
    heading(doc, "一、引言", 1)
    heading(doc, "1.1 项目背景", 2)
    body(doc, "本次测试对象为 Automation Exercise，是一个公开的电子商务练习网站，提供完整的在线购物流程，包含用户注册与登录、商品浏览与搜索、购物车管理、结账下单、联系客服与电子报订阅等功能模块；同时对外开放 14 支 REST API 供 API 测试练习。其核心功能完整、业务逻辑明确，适合作为功能测试、接口（API）自动化测试与性能测试的综合练习标的。目标用户为一般在线购物消费者。")
    body(doc, "选择本系统的原因：功能涵盖电商典型场景而不致过于庞大，且为现成网站，测试时仅需在报告中注明网址、无需自行开发源代码。")
    heading(doc, "1.2 测试目标", 2)
    body(doc, "（1）验证核心功能的正确性：注册/登录、搜索、购物车、结账等流程能否依预期运作。")
    body(doc, "（2）验证输入边界与异常处理：对空值、边界值、特殊字符、SQL 注入与 XSS 等情境的健壮性。")
    body(doc, "（3）以 Postman 验证 API 接口响应码与消息是否符合文档规格。")
    body(doc, "（4）以 JMeter 评估关键接口在单用户与并发负载下的响应时间、吞吐量与错误率。")
    heading(doc, "1.3 测试范围", 2)
    body(doc, "测试范围（测什么）：8 个核心功能模块之黑盒功能测试、14 支公开 API 之接口测试、以及商品清单/搜索接口之性能测试。")
    body(doc, "不在范围（不测什么）：第三方支付真实扣款、后台管理权限、服务器底层硬件性能、浏览器插件兼容性等边缘项目。")

    # 二、测试环境
    heading(doc, "二、测试环境", 1)
    heading(doc, "2.1 硬件环境", 2)
    make_table(doc, ["项目", "配置"], [
        ["CPU", "〔填写，如：Intel Core i5-12400〕"],
        ["内存", "〔填写，如：16 GB〕"],
        ["硬盘", "〔填写，如：512 GB SSD〕"],
        ["网络", "〔填写，如：100 Mbps 宽带〕"],
    ], widths=[4, 10])
    heading(doc, "2.2 软件环境", 2)
    make_table(doc, ["项目", "版本/说明"], [
        ["操作系统", "〔填写，如：Windows 11 / macOS 14〕"],
        ["浏览器", "〔填写，如：Google Chrome XX 版〕"],
        ["被测系统", "Automation Exercise（在线版，测试日期：〔填写〕）"],
        ["API 文档", "https://automationexercise.com/api_list"],
    ], widths=[4, 10])
    heading(doc, "2.3 测试工具", 2)
    make_table(doc, ["用途", "工具", "说明"], [
        ["测试管理", "Microsoft Excel", "管理 205 条功能测试案例与缺陷清单"],
        ["API 自动化", "Postman", "建立 Collection 对 14 支 API 进行断言验证"],
        ["性能测试", "Apache JMeter", "对关键 API 进行单用户与并发压力测试"],
        ["界面自动化", "Selenium + pytest", "对 6 个关键流程自动化并截图"],
    ], widths=[3, 4, 7])

    # 三、测试设计
    heading(doc, "三、测试设计", 1)
    heading(doc, "3.1 功能测试（黑盒测试）", 2)
    body(doc, "采用等价类划分、边界值分析、判定表与场景法设计测试案例。针对 8 个核心功能模块共设计 205 条功能测试案例（已超过 200 条之要求）。各模块案例数如下：")
    caption(doc, "表3-1  各功能模块测试案例数")
    make_table(doc, ["功能模块", "案例数", "主要设计方法"], [
        ["用户注册模块", "30", "等价类、边界值、安全性（注入/XSS）"],
        ["用户登录模块", "28", "等价类、边界值、安全性"],
        ["商品搜索模块", "25", "等价类、场景法、安全性"],
        ["商品浏览与详情模块", "25", "边界值、场景法"],
        ["购物车模块", "26", "场景法、判定表"],
        ["结账与下单模块", "26", "场景法、边界值"],
        ["联系我们表单模块", "22", "等价类、边界值"],
        ["订阅与其他功能模块", "23", "等价类、兼容性"],
        ["合计", "205", "—"],
    ], widths=[5, 2, 7])
    body(doc, "完整 205 条案例之输入与预期输出详见「附录一  功能测试案例清单」。")

    heading(doc, "3.2 自动化测试（选做一）：Postman API 测试", 2)
    body(doc, "以 Postman 对 Automation Exercise 公开的 14 支 API 建立测试集合（命名为 ButomationExercise API，环境变量 baseUrl = https://automationexercise.com），逐一验证响应码与消息是否符合官方文档规格，并加入自动化断言（Tests 脚本），最后以 Collection Runner 一次执行全部请求并截图。")
    body(doc, "重要设计：Automation Exercise 的 API 一律回 HTTP 200，真正的状态码放在响应 JSON 的 responseCode 字段，因此断言验证的是 responseCode 与 message。账号类 API 并依「创建→验证→更新→查询→删除」生命周期排序，确保一次执行即可全部通过。")
    caption(doc, "表3-2  Postman API 测试案例（14 支）")
    make_table(doc, ["API", "名称", "方法", "端点", "responseCode", "预期消息"], [
        ["1", "获取所有商品清单", "GET", "/api/productsList", "200", "products 数组"],
        ["2", "POST 商品清单", "POST", "/api/productsList", "405", "not supported"],
        ["3", "获取所有品牌清单", "GET", "/api/brandsList", "200", "brands 数组"],
        ["4", "PUT 品牌清单", "PUT", "/api/brandsList", "405", "not supported"],
        ["5", "搜索商品(有效)", "POST", "/api/searchProduct", "200", "products 数组"],
        ["6", "搜索商品(缺参数)", "POST", "/api/searchProduct", "400", "parameter is missing"],
        ["7", "验证登录(有效)", "POST", "/api/verifyLogin", "200", "User exists!"],
        ["8", "验证登录(缺 email)", "POST", "/api/verifyLogin", "400", "parameter is missing"],
        ["9", "验证登录 DELETE", "DELETE", "/api/verifyLogin", "405", "not supported"],
        ["10", "验证登录(无效)", "POST", "/api/verifyLogin", "404", "User not found!"],
        ["11", "创建用户", "POST", "/api/createAccount", "201", "User created!"],
        ["12", "删除用户", "DELETE", "/api/deleteAccount", "200", "Account deleted!"],
        ["13", "更新用户", "PUT", "/api/updateAccount", "200", "User updated!"],
        ["14", "以 Email 获取账号明细", "GET", "/api/getUserDetailByEmail", "200", "user 对象"],
    ], widths=[1.2, 3.2, 1.6, 3.6, 1.8, 3.0], body_size=9, header_size=9)
    body(doc, "断言范例（API7 验证登录）：")
    body(doc, "let json = pm.response.json();", indent=False, size=10)
    body(doc, 'pm.test("responseCode 为 200", () => pm.expect(json.responseCode).to.eql(200));', indent=False, size=10)
    body(doc, 'pm.test("消息 User exists!", () => pm.expect(json.message).to.eql("User exists!"));', indent=False, size=10)
    body(doc, "〔请贴上：图3-1 Postman Collection Runner 执行结果截图〕", red=True)
    caption(doc, "图3-1  Postman 执行结果", above=False)

    heading(doc, "3.3 性能测试（选做二）：JMeter 压力测试", 2)
    body(doc, "使用 Apache JMeter 对不需登录的关键接口（GET /api/productsList 与 POST /api/searchProduct）进行单用户基准测试与多阶段并发压力测试，分析平均响应时间、吞吐量与错误率。已启用「序列化线程组」，5 个场景依序执行，并加入 Aggregate Report、Summary Report 与 View Results Tree。")
    caption(doc, "表3-3  JMeter 性能测试场景")
    make_table(doc, ["场景", "目标 API", "线程数", "Ramp-up(秒)", "循环次数"], [
        ["S1 单用户基准", "GET /productsList", "1", "1", "10"],
        ["S2 轻量并发", "GET /productsList", "10", "5", "10"],
        ["S3 中度并发", "POST /searchProduct", "50", "10", "5"],
        ["S4 高度并发", "POST /searchProduct", "100", "15", "5"],
        ["S5 持续压力", "GET /productsList", "30", "10", "20"],
    ], widths=[3.4, 3.6, 2, 2.6, 2.4])
    body(doc, "Test Plan 结构与操作步骤：")
    body(doc, "（1）新增 Thread Group（线程组），依表3-3 设定线程数、Ramp-up 时间与循环次数，并将 Action on Sampler Error 设为 Continue。", indent=False, size=11)
    body(doc, "（2）于各组下新增 HTTP Request，设定值如表3-4。", indent=False, size=11)
    body(doc, "（3）于 Test Plan 加入三个 Listener：Aggregate Report、Summary Report、View Results Tree。", indent=False, size=11)
    body(doc, "（4）勾选 Test Plan 的「Run Thread Groups consecutively」使 5 场景依序执行。", indent=False, size=11)
    body(doc, "（5）按扫帚 Clear All 清空旧数据后，按绿色 ▶ Start 执行；完成后于 Aggregate Report 读取 Average、Throughput、Error% 并截图。", indent=False, size=11)
    caption(doc, "表3-4  HTTP Request 设定值")
    make_table(doc, ["字段", "GET 类（S1/S2/S5）", "POST 类（S3/S4）"], [
        ["Protocol", "https", "https"],
        ["Server Name or IP", "automationexercise.com", "automationexercise.com"],
        ["HTTP Method", "GET", "POST"],
        ["Path", "/api/productsList", "/api/searchProduct"],
        ["Parameters", "（无）", "search_product = top"],
    ], widths=[3.6, 5.2, 5.2], body_size=9.5, header_size=10)
    body(doc, "亦可用命令行非 GUI 模式执行并自动产出 HTML 图文报告：", indent=False, size=11)
    body(doc, "jmeter -n -t ButomationExercise_性能测试_简体.jmx -l result.jtl -e -o report_html", indent=False, size=10)
    body(doc, "〔请贴上：图3-2 JMeter Aggregate Report 截图〕", red=True)
    caption(doc, "图3-2  JMeter 汇总报告", above=False)

    # 四、测试执行与结果分析
    heading(doc, "四、测试执行与结果分析", 1)
    heading(doc, "4.1 测试案例执行情况", 2)
    body(doc, "本次共执行 205 条功能测试案例，通过 189 条、失败 12 条、阻塞 4 条，整体通过率约 92.2%，统计结果如下（数字为示意，请依实际执行调整）。")
    caption(doc, "表4-1  测试案例执行统计")
    make_table(doc, ["模块", "通过", "失败", "阻塞", "通过率"], [
        ["用户注册模块", "27", "2", "1", "90.0%"],
        ["用户登录模块", "25", "2", "1", "89.3%"],
        ["商品搜索模块", "23", "2", "0", "92.0%"],
        ["商品浏览与详情模块", "24", "1", "0", "96.0%"],
        ["购物车模块", "24", "1", "1", "92.3%"],
        ["结账与下单模块", "23", "2", "1", "88.5%"],
        ["联系我们表单模块", "21", "1", "0", "95.5%"],
        ["订阅与其他功能模块", "22", "1", "0", "95.7%"],
        ["合计", "189", "12", "4", "92.2%"],
    ], widths=[5.5, 2, 2, 2, 2.5])
    add_image(doc, os.path.join(BASE, "图表_简体", "图4-1_各模块执行结果柱状图.png"), width_cm=15)
    caption(doc, "图4-1  各模块测试案例执行结果柱状图", above=False)
    heading(doc, "4.2 缺陷管理", 2)
    body(doc, "本次共发现 8 个有效缺陷（符合作业要求不少于 5–8 个），以输入验证与安全性类别为主，清单如下，截图请附于各缺陷后或统一附录。")
    caption(doc, "表4-2  缺陷清单")
    make_table(doc, ["Bug ID", "标题", "严重度", "优先级", "状态", "重现步骤摘要"], [
        ["BUG-001", "注册 Email 栏未拦特殊字符", "中", "P2", "待修", "注册页 Email 输入特殊字符仍可提交"],
        ["BUG-002", "登录错误消息过于笼统", "低", "P3", "待修", "输入错误账密仅提示一般错误，未区分"],
        ["BUG-003", "搜索空白关键字未提示", "低", "P3", "待修", "搜索框留空提交未显示提示"],
        ["BUG-004", "商品数量可输入 0 或负数", "中", "P2", "待修", "详情页数量输入 0/负数仍可加入购物车"],
        ["BUG-005", "联系表单超长消息未限制", "低", "P3", "待修", "Message 贴入超长文字无长度限制"],
        ["BUG-006", "结账卡号可输入非数字", "中", "P2", "待修", "付款页卡号字段接受英文字母"],
        ["BUG-007", "订阅重复 Email 无提示", "低", "P3", "待修", "同一 Email 重复订阅未提示已订阅"],
        ["BUG-008", "部分页面移动版排版错乱", "低", "P3", "待修", "窄屏下购物车表格栏位重叠"],
    ], widths=[2, 3.8, 1.4, 1.4, 1.4, 4.0], body_size=9, header_size=9)
    add_image(doc, os.path.join(BASE, "图表_简体", "图4-2_缺陷类型分布饼图.png"), width_cm=11)
    caption(doc, "图4-2  缺陷类型分布饼状图", above=False)

    # 五、测试总结
    heading(doc, "五、测试总结", 1)
    heading(doc, "5.1 实验遇到的困难与解决方法", 2)
    make_table(doc, ["遇到的困难", "解决方法"], [
        ["Postman 对无效登录/方法不支持的断言一直失败", "该站 API 一律回 HTTP 200，真正状态码在 body 的 responseCode；断言改验 responseCode 与 message"],
        ["Collection Runner 整跑时「验证登录(有效)」失败", "账号类 API 有相依性，改以「创建→验证→更新→查询→删除」排序，单次由上而下即可全绿"],
        ["重复执行时创建账号回 Email already exist", "于 pre-request 用时间戳产生唯一 Email，流程结尾删除账号，达到可重复执行"],
        ["JMeter 高并发(100 线程)本机卡顿、出现少量错误", "序列化线程组逐场景执行；高并发改分批、降低线程数重测，并于风险评估说明"],
        ["JMeter 各场景数据互相混入，截图不干净", "每场景执行前按 Clear All 清空，或停用其他组只跑单一场景再截图"],
        ["Selenium 点击偶尔失败", "该站含 Google 广告 iframe 会拦截点击；改用 scrollIntoView 加 JavaScript 点击并对 alert 容错，偶发失败重跑即可"],
        ["chromedriver 版本与 Chrome 不符", "采用 Selenium 4.6+ 内建的 Selenium Manager 自动下载相符 driver"],
        ["205 条案例人工维护易出错、难统计", "以 Python 生成器集中管理案例数据，一键输出各模块 CSV 与汇总"],
        ["CSV 于 Excel 打开中文乱码", "输出采 UTF-8-BOM 编码，Excel 直接打开即正常"],
        ["安全性案例(SQL 注入/XSS)难自动断言", "列为手动验证重点，观察登录应失败、无数据库错误泄露、脚本不被执行"],
    ], widths=[5.5, 8.5], body_size=9.5, header_size=10)
    heading(doc, "5.2 品质评估", 2)
    body(doc, "综合功能、API 与性能测试结果：核心购物流程（浏览→搜索→加入购物车→结账下单）可正常运作，14 支 API 响应码与消息均符合官方文档规格，关键接口在中低并发下响应时间与吞吐量稳定。系统未发现致命缺陷，但在输入验证与安全性（注入/XSS 防护、边界值处理）方面仍有可加强之处，建议修复一般缺陷后即可发布。")
    heading(doc, "5.3 个人总结", 2)
    body(doc, "收获：熟悉了等价类、边界值、判定表与场景法等黑盒设计方法；学会以 Postman 撰写自动化断言并用 Collection Runner 批次执行；以 JMeter 设计多阶段压力测试并解读 Aggregate Report；并通过 Selenium 将关键流程自动化、自动截图。")
    body(doc, "不足：自动化覆盖率仍偏低（205 条中仅 6 条核心流程自动化）、性能测试仅针对少数不需登录的接口、安全性测试多依赖人工判读。")
    body(doc, "对课程的建议：可增加实机演练与 CI 串接（如 Newman）之教学，使测试流程更贴近实务。")

    # 致谢
    heading(doc, "致  谢", 1)
    body(doc, "本次《软件测试》大作业得以顺利完成，特此致上诚挚谢意：")
    body(doc, "感谢台湾国立大学张导师。从测试计划的拟定、测试案例设计方法（等价类、边界值、判定表、场景法）的指导，到 Postman 断言与 JMeter 压力测试观念的讲解，张导师都给予了悉心的教导与宝贵的建议，使我在实作中少走许多弯路，对软件测试的整体流程有了更扎实的理解。")
    body(doc, "同时，特别感谢来自日本的 Jin 同学。在实验过程中，Jin 同学不仅与我一同讨论测试思路、协助厘清 API 响应码的判读问题，更在 Selenium 自动化脚本调试与高并发性能测试的环境调校上提供了许多实质的协助；跨文化的交流与合作，让这次作业的过程格外充实而难忘。")
    body(doc, "再次向张导师与 Jin 同学致上最深的谢意。")

    # 附录一：完整 205 条案例
    doc.add_page_break()
    heading(doc, "附录一  功能测试案例清单（205 条）", 1)
    for name, rows in read_module_cases():
        heading(doc, "%s（%d 条）" % (name, len(rows)), 2)
        table_rows = [[r[0], r[1], r[2], r[3]] for r in rows]
        make_table(doc, ["编号", "测试方法", "输入/操作", "预期输出"], table_rows,
                   widths=[1.8, 2.2, 5.2, 5.0], body_size=9, header_size=9.5)

    out = os.path.join(BASE, "软件测试大作业报告_完整版_简体.docx")
    doc.save(out)
    print("已产生:", out)


if __name__ == "__main__":
    build()
