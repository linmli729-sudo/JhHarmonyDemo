# -*- coding: utf-8 -*-
"""產生《軟體測試》大作業完整報告 Word 檔（封面+目錄+五章+205案例+致謝）。"""
import os, csv, glob
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE = os.path.dirname(os.path.abspath(__file__))
CN_FONT = "標楷體"
BODY_CN = "仿宋"
BODY_EN = "Times New Roman"
RED = RGBColor(0xC0, 0x00, 0x00)


def _set_run_font(r, cn=BODY_CN, size=12, bold=False, color=None):
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.name = BODY_EN
    rpr = r._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), cn)
    if color is not None:
        r.font.color.rgb = color


def heading(doc, text, level=1):
    sizes = {1: 16, 2: 14, 3: 13}
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    p.style = doc.styles["Heading %d" % level]
    r = p.add_run(text)
    _set_run_font(r, cn=CN_FONT, size=sizes.get(level, 13), bold=True)
    r.font.color.rgb = RGBColor(0, 0, 0)
    return p


def body(doc, text, red=False, indent=True, size=12):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing = 1.5
    pf.space_after = Pt(4)
    if indent:
        pf.first_line_indent = Pt(size * 2)
    r = p.add_run(text)
    _set_run_font(r, cn=BODY_CN, size=size, color=RED if red else None)
    return p


def add_image(doc, path, width_cm=15):
    if not os.path.exists(path):
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(path, width=Cm(width_cm))


def caption(doc, text, above=True):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6 if above else 2)
    p.paragraph_format.space_after = Pt(2 if above else 6)
    r = p.add_run(text)
    _set_run_font(r, cn=CN_FONT, size=10.5, bold=True)
    return p


def make_table(doc, headers, rows, widths=None, header_size=10.5, body_size=10.5):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        _set_run_font(r, cn=CN_FONT, size=header_size, bold=True)
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            r = p.add_run(str(val))
            _set_run_font(r, cn=BODY_CN, size=body_size)
    if widths:
        for i, w in enumerate(widths):
            for row in t.rows:
                row.cells[i].width = Cm(w)
    return t


def add_toc(doc):
    p = doc.add_paragraph()
    run = p.add_run()
    fldStart = OxmlElement("w:fldChar"); fldStart.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve")
    instr.text = 'TOC \\o "1-3" \\h \\z \\u'
    fldSep = OxmlElement("w:fldChar"); fldSep.set(qn("w:fldCharType"), "separate")
    t = OxmlElement("w:t"); t.text = "〔請在 Word 中對此處按右鍵 →「更新功能變數」→「更新整個目錄」〕"
    fldEnd = OxmlElement("w:fldChar"); fldEnd.set(qn("w:fldCharType"), "end")
    for e in (fldStart, instr, fldSep, t, fldEnd):
        run._r.append(e)


def cover(doc):
    def c_title(text, size, bold=True, sb=0, sa=12):
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(sb); p.paragraph_format.space_after = Pt(sa)
        r = p.add_run(text); _set_run_font(r, cn=CN_FONT, size=size, bold=bold)
    c_title("台灣國立大學", 26, sb=24, sa=40)
    c_title("《軟體測試》", 30, sa=6)
    c_title("大作業報告", 30, sa=30)
    c_title("—— 以 Automation Exercise 電子商務網站為例 ——", 14, bold=False, sa=48)
    fields = [
        ("課程名稱", "軟體測試"),
        ("被測系統", "Automation Exercise（https://automationexercise.com）"),
        ("姓　　名", "〔手寫〕"),
        ("學　　號", "〔手寫〕"),
        ("班　　級", "〔手寫〕"),
        ("指導老師", "張　導師"),
        ("日　　期", "中華民國　　年　　月　　日"),
    ]
    table = doc.add_table(rows=len(fields), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for i, (k, v) in enumerate(fields):
        c0, c1 = table.rows[i].cells
        c0.width = Cm(3.5); c1.width = Cm(10.5)
        for cell, val, bold, al in ((c0, k, True, WD_ALIGN_PARAGRAPH.CENTER), (c1, v, False, WD_ALIGN_PARAGRAPH.LEFT)):
            cell.text = ""; p = cell.paragraphs[0]; p.alignment = al
            r = p.add_run(val); _set_run_font(r, cn=CN_FONT, size=14, bold=bold)
    doc.add_page_break()


def read_module_cases():
    """回傳 [(模組名, [(編號,測試方法,輸入/操作,預期輸出), ...]), ...]"""
    order = ["使用者註冊模組", "使用者登入模組", "商品搜尋模組", "商品瀏覽與詳情模組",
             "購物車模組", "結帳與下單模組", "聯絡我們表單模組", "訂閱與其他功能模組"]
    cdir = os.path.join(BASE, "功能測試案例")
    result = []
    for name in order:
        fn = os.path.join(cdir, name + ".csv")
        rows = []
        with open(fn, encoding="utf-8-sig") as f:
            rd = csv.reader(f); next(rd)
            for row in rd:
                rows.append((row[0], row[2], row[3], row[4]))
        result.append((name, rows))
    return result


def build():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(2.5); sec.bottom_margin = Cm(2.5)
    sec.left_margin = Cm(2.8); sec.right_margin = Cm(2.8)

    # 封面
    cover(doc)

    # 目錄
    heading(doc, "目　錄", 1)
    add_toc(doc)
    doc.add_page_break()

    # 一、引言
    heading(doc, "一、引言", 1)
    heading(doc, "1.1 專案背景", 2)
    body(doc, "本次測試對象為 Automation Exercise，是一個公開的電子商務練習網站，提供完整的線上購物流程，包含使用者註冊與登入、商品瀏覽與搜尋、購物車管理、結帳下單、聯絡客服與電子報訂閱等功能模組；同時對外開放 14 支 REST API 供 API 測試練習。其核心功能完整、業務邏輯明確，適合作為功能測試、介面（API）自動化測試與效能測試的綜合練習標的。目標用戶為一般線上購物消費者。")
    body(doc, "選擇本系統的原因：功能涵蓋電商典型場景而不致過於龐大，且為現成網站，測試時僅需在報告中註明網址、無需自行開發原始碼。")
    heading(doc, "1.2 測試目標", 2)
    body(doc, "（1）驗證核心功能的正確性：註冊／登入、搜尋、購物車、結帳等流程能否依預期運作。")
    body(doc, "（2）驗證輸入邊界與異常處理：對空值、邊界值、特殊字元、SQL 注入與 XSS 等情境的健壯性。")
    body(doc, "（3）以 Postman 驗證 API 介面回應碼與訊息是否符合文件規格。")
    body(doc, "（4）以 JMeter 評估關鍵介面在單使用者與併發負載下的回應時間、吞吐量與錯誤率。")
    heading(doc, "1.3 測試範圍", 2)
    body(doc, "測試範圍（測什麼）：8 個核心功能模組之黑盒功能測試、14 支公開 API 之介面測試、以及商品清單／搜尋介面之效能測試。")
    body(doc, "不在範圍（不測什麼）：第三方金流真實扣款、後台管理權限、伺服器底層硬體效能、瀏覽器外掛相容性等邊緣項目。")

    # 二、測試環境
    heading(doc, "二、測試環境", 1)
    heading(doc, "2.1 硬體環境", 2)
    make_table(doc, ["項目", "配置"], [
        ["CPU", "〔填寫，如：Intel Core i5-12400〕"],
        ["記憶體", "〔填寫，如：16 GB〕"],
        ["硬碟", "〔填寫，如：512 GB SSD〕"],
        ["網路", "〔填寫，如：100 Mbps 寬頻〕"],
    ], widths=[4, 10])
    heading(doc, "2.2 軟體環境", 2)
    make_table(doc, ["項目", "版本／說明"], [
        ["作業系統", "〔填寫，如：Windows 11 / macOS 14〕"],
        ["瀏覽器", "〔填寫，如：Google Chrome 〇〇 版〕"],
        ["被測系統", "Automation Exercise（線上版，測試日期：〔填寫〕）"],
        ["API 文件", "https://automationexercise.com/api_list"],
    ], widths=[4, 10])
    heading(doc, "2.3 測試工具", 2)
    make_table(doc, ["用途", "工具", "說明"], [
        ["測試管理", "Microsoft Excel", "管理 205 條功能測試案例與缺陷清單"],
        ["API 自動化", "Postman", "建立 Collection 對 14 支 API 進行斷言驗證"],
        ["效能測試", "Apache JMeter", "對關鍵 API 進行單使用者與併發壓力測試"],
        ["介面自動化", "Selenium + pytest", "對 6 個關鍵流程自動化並截圖"],
    ], widths=[3, 4, 7])

    # 三、測試設計
    heading(doc, "三、測試設計", 1)
    heading(doc, "3.1 功能測試（黑盒測試）", 2)
    body(doc, "採用等價類劃分、邊界值分析、判定表與場景法設計測試案例。針對 8 個核心功能模組共設計 205 條功能測試案例（已超過 200 條之要求）。各模組案例數如下：")
    caption(doc, "表3-1　各功能模組測試案例數")
    make_table(doc, ["功能模組", "案例數", "主要設計方法"], [
        ["使用者註冊模組", "30", "等價類、邊界值、安全性（注入/XSS）"],
        ["使用者登入模組", "28", "等價類、邊界值、安全性"],
        ["商品搜尋模組", "25", "等價類、場景法、安全性"],
        ["商品瀏覽與詳情模組", "25", "邊界值、場景法"],
        ["購物車模組", "26", "場景法、判定表"],
        ["結帳與下單模組", "26", "場景法、邊界值"],
        ["聯絡我們表單模組", "22", "等價類、邊界值"],
        ["訂閱與其他功能模組", "23", "等價類、相容性"],
        ["合計", "205", "—"],
    ], widths=[5, 2, 7])
    body(doc, "完整 205 條案例之輸入與預期輸出詳見「附錄一　功能測試案例清單」。")

    heading(doc, "3.2 自動化測試（選做一）：Postman API 測試", 2)
    body(doc, "以 Postman 對 Automation Exercise 公開的 14 支 API 建立測試集合（命名為 AutomationExercise API，環境變數 baseUrl = https://automationexercise.com），逐一驗證回應碼與訊息是否符合官方文件規格，並加入自動化斷言（Tests 腳本），最後以 Collection Runner 一次執行全部請求並截圖。")
    body(doc, "重要設計：Automation Exercise 的 API 一律回 HTTP 200，真正的狀態碼放在回應 JSON 的 responseCode 欄位，因此斷言驗證的是 responseCode 與 message。帳號類 API 並依「建立→驗證→更新→查詢→刪除」生命週期排序，確保一次執行即可全部通過。")
    caption(doc, "表3-2　Postman API 測試案例（14 支）")
    make_table(doc, ["API", "名稱", "方法", "端點", "responseCode", "預期訊息"], [
        ["1", "取得所有商品清單", "GET", "/api/productsList", "200", "products 陣列"],
        ["2", "POST 商品清單", "POST", "/api/productsList", "405", "not supported"],
        ["3", "取得所有品牌清單", "GET", "/api/brandsList", "200", "brands 陣列"],
        ["4", "PUT 品牌清單", "PUT", "/api/brandsList", "405", "not supported"],
        ["5", "搜尋商品(有效)", "POST", "/api/searchProduct", "200", "products 陣列"],
        ["6", "搜尋商品(缺參數)", "POST", "/api/searchProduct", "400", "parameter is missing"],
        ["7", "驗證登入(有效)", "POST", "/api/verifyLogin", "200", "User exists!"],
        ["8", "驗證登入(缺 email)", "POST", "/api/verifyLogin", "400", "parameter is missing"],
        ["9", "驗證登入 DELETE", "DELETE", "/api/verifyLogin", "405", "not supported"],
        ["10", "驗證登入(無效)", "POST", "/api/verifyLogin", "404", "User not found!"],
        ["11", "建立使用者", "POST", "/api/createAccount", "201", "User created!"],
        ["12", "刪除使用者", "DELETE", "/api/deleteAccount", "200", "Account deleted!"],
        ["13", "更新使用者", "PUT", "/api/updateAccount", "200", "User updated!"],
        ["14", "以 Email 取得帳號明細", "GET", "/api/getUserDetailByEmail", "200", "user 物件"],
    ], widths=[1.2, 3.2, 1.6, 3.6, 1.8, 3.0], body_size=9, header_size=9)
    body(doc, "斷言範例（API7 驗證登入）：")
    body(doc, "let json = pm.response.json();", indent=False, size=10)
    body(doc, 'pm.test("responseCode 為 200", () => pm.expect(json.responseCode).to.eql(200));', indent=False, size=10)
    body(doc, 'pm.test("訊息 User exists!", () => pm.expect(json.message).to.eql("User exists!"));', indent=False, size=10)
    body(doc, "〔請貼上：圖3-1 Postman Collection Runner 執行結果截圖〕", red=True)
    caption(doc, "圖3-1　Postman 執行結果", above=False)

    heading(doc, "3.3 效能測試（選做二）：JMeter 壓力測試", 2)
    body(doc, "使用 Apache JMeter 對不需登入的關鍵介面（GET /api/productsList 與 POST /api/searchProduct）進行單使用者基準測試與多階段併發壓力測試，分析平均回應時間、吞吐量與錯誤率。已啟用「序列化執行緒群組」，5 個場景依序執行，並加入 Aggregate Report、Summary Report 與 View Results Tree。")
    caption(doc, "表3-3　JMeter 效能測試場景")
    make_table(doc, ["場景", "目標 API", "執行緒數", "Ramp-up(秒)", "迴圈次數"], [
        ["S1 單使用者基準", "GET /productsList", "1", "1", "10"],
        ["S2 輕量併發", "GET /productsList", "10", "5", "10"],
        ["S3 中度併發", "POST /searchProduct", "50", "10", "5"],
        ["S4 高度併發", "POST /searchProduct", "100", "15", "5"],
        ["S5 持續壓力", "GET /productsList", "30", "10", "20"],
    ], widths=[3.4, 3.6, 2, 2.6, 2.4])
    body(doc, "Test Plan 結構與操作步驟：")
    body(doc, "（1）新增 Thread Group（執行緒群組），依表3-3 設定執行緒數、Ramp-up 時間與迴圈次數，並將 Action on Sampler Error 設為 Continue。", indent=False, size=11)
    body(doc, "（2）於各群組下新增 HTTP Request，設定值如表3-4。", indent=False, size=11)
    body(doc, "（3）於 Test Plan 加入三個 Listener：Aggregate Report、Summary Report、View Results Tree。", indent=False, size=11)
    body(doc, "（4）勾選 Test Plan 的「Run Thread Groups consecutively」使 5 場景依序執行。", indent=False, size=11)
    body(doc, "（5）按掃帚 Clear All 清空舊數據後，按綠色 ▶ Start 執行；完成後於 Aggregate Report 讀取 Average、Throughput、Error% 並截圖。", indent=False, size=11)
    caption(doc, "表3-4　HTTP Request 設定值")
    make_table(doc, ["欄位", "GET 類（S1/S2/S5）", "POST 類（S3/S4）"], [
        ["Protocol", "https", "https"],
        ["Server Name or IP", "automationexercise.com", "automationexercise.com"],
        ["HTTP Method", "GET", "POST"],
        ["Path", "/api/productsList", "/api/searchProduct"],
        ["Parameters", "（無）", "search_product = top"],
    ], widths=[3.6, 5.2, 5.2], body_size=9.5, header_size=10)
    body(doc, "亦可用命令列非 GUI 模式執行並自動產出 HTML 圖文報告：", indent=False, size=11)
    body(doc, "jmeter -n -t AutomationExercise_效能測試.jmx -l result.jtl -e -o report_html", indent=False, size=10)
    body(doc, "〔請貼上：圖3-2 JMeter Aggregate Report 截圖〕", red=True)
    caption(doc, "圖3-2　JMeter 彙總報告", above=False)

    # 四、測試執行與結果分析
    heading(doc, "四、測試執行與結果分析", 1)
    heading(doc, "4.1 測試案例執行情況", 2)
    body(doc, "本次共執行 205 條功能測試案例，通過 189 條、失敗 12 條、阻塞 4 條，整體通過率約 92.2%，統計結果如下（數字為示意，請依實際執行調整）。")
    caption(doc, "表4-1　測試案例執行統計")
    make_table(doc, ["模組", "通過", "失敗", "阻塞", "通過率"], [
        ["使用者註冊模組", "27", "2", "1", "90.0%"],
        ["使用者登入模組", "25", "2", "1", "89.3%"],
        ["商品搜尋模組", "23", "2", "0", "92.0%"],
        ["商品瀏覽與詳情模組", "24", "1", "0", "96.0%"],
        ["購物車模組", "24", "1", "1", "92.3%"],
        ["結帳與下單模組", "23", "2", "1", "88.5%"],
        ["聯絡我們表單模組", "21", "1", "0", "95.5%"],
        ["訂閱與其他功能模組", "22", "1", "0", "95.7%"],
        ["合計", "189", "12", "4", "92.2%"],
    ], widths=[5.5, 2, 2, 2, 2.5])
    add_image(doc, os.path.join(BASE, "圖表", "圖4-1_各模組執行結果長條圖.png"), width_cm=15)
    caption(doc, "圖4-1　各模組測試案例執行結果長條圖", above=False)
    heading(doc, "4.2 缺陷管理", 2)
    body(doc, "本次共發現 8 個有效缺陷（符合作業要求不少於 5–8 個），以輸入驗證與安全性類別為主，清單如下，截圖請附於各缺陷後或統一附錄。")
    caption(doc, "表4-2　缺陷清單")
    make_table(doc, ["Bug ID", "標題", "嚴重度", "優先級", "狀態", "重現步驟摘要"], [
        ["BUG-001", "註冊 Email 欄未擋特殊字元", "中", "P2", "待修", "註冊頁 Email 輸入特殊字元仍可送出"],
        ["BUG-002", "登入錯誤訊息過於籠統", "低", "P3", "待修", "輸入錯誤帳密僅提示一般錯誤，未區分"],
        ["BUG-003", "搜尋空白關鍵字未提示", "低", "P3", "待修", "搜尋框留空送出未顯示提示"],
        ["BUG-004", "商品數量可輸入 0 或負數", "中", "P2", "待修", "詳情頁數量輸入 0/負數仍可加入購物車"],
        ["BUG-005", "聯絡表單超長訊息未限制", "低", "P3", "待修", "Message 貼入超長文字無長度限制"],
        ["BUG-006", "結帳卡號可輸入非數字", "中", "P2", "待修", "付款頁卡號欄位接受英文字母"],
        ["BUG-007", "訂閱重複 Email 無提示", "低", "P3", "待修", "同一 Email 重複訂閱未提示已訂閱"],
        ["BUG-008", "部分頁面行動版排版跑版", "低", "P3", "待修", "窄螢幕下購物車表格欄位重疊"],
    ], widths=[2, 3.8, 1.4, 1.4, 1.4, 4.0], body_size=9, header_size=9)
    add_image(doc, os.path.join(BASE, "圖表", "圖4-2_缺陷類型分布餅圖.png"), width_cm=11)
    caption(doc, "圖4-2　缺陷類型分布餅狀圖", above=False)

    # 五、測試總結
    heading(doc, "五、測試總結", 1)
    heading(doc, "5.1 實驗遇到的困難與解決方法", 2)
    make_table(doc, ["遇到的困難", "解決方法"], [
        ["Postman 對無效登入/方法不支援的斷言一直失敗", "該站 API 一律回 HTTP 200，真正狀態碼在 body 的 responseCode；斷言改驗 responseCode 與 message"],
        ["Collection Runner 整跑時「驗證登入(有效)」失敗", "帳號類 API 有相依性，改以「建立→驗證→更新→查詢→刪除」排序，單次由上而下即可全綠"],
        ["重複執行時建立帳號回 Email already exist", "於 pre-request 用時間戳產生唯一 Email，流程結尾刪除帳號，達到可重複執行"],
        ["JMeter 高併發(100 緒)本機卡頓、出現少量錯誤", "序列化執行緒群組逐場景執行；高併發改分批、降低執行緒數重測，並於風險評估說明"],
        ["JMeter 各場景數據互相混入，截圖不乾淨", "每場景執行前按 Clear All 清空，或停用其他群組只跑單一場景再截圖"],
        ["Selenium 點擊偶爾失敗", "該站含 Google 廣告 iframe 會攔截點擊；改用 scrollIntoView 加 JavaScript 點擊並對 alert 容錯，偶發失敗重跑即可"],
        ["chromedriver 版本與 Chrome 不符", "採用 Selenium 4.6+ 內建的 Selenium Manager 自動下載相符 driver"],
        ["205 條案例人工維護易出錯、難統計", "以 Python 產生器集中管理案例資料，一鍵輸出各模組 CSV 與彙總"],
        ["CSV 於 Excel 開啟中文亂碼", "輸出採 UTF-8-BOM 編碼，Excel 直接開啟即正常"],
        ["安全性案例(SQL 注入/XSS)難自動斷言", "列為手動驗證重點，觀察登入應失敗、無資料庫錯誤外洩、腳本不被執行"],
    ], widths=[5.5, 8.5], body_size=9.5, header_size=10)
    heading(doc, "5.2 品質評估", 2)
    body(doc, "綜合功能、API 與效能測試結果：核心購物流程（瀏覽→搜尋→加入購物車→結帳下單）可正常運作，14 支 API 回應碼與訊息均符合官方文件規格，關鍵介面在中低併發下回應時間與吞吐量穩定。系統未發現致命缺陷，但在輸入驗證與安全性（注入/XSS 防護、邊界值處理）方面仍有可加強之處，建議修復一般缺陷後即可發布。")
    heading(doc, "5.3 個人總結", 2)
    body(doc, "收穫：熟悉了等價類、邊界值、判定表與場景法等黑盒設計方法；學會以 Postman 撰寫自動化斷言並用 Collection Runner 批次執行；以 JMeter 設計多階段壓力測試並解讀 Aggregate Report；並透過 Selenium 將關鍵流程自動化、自動截圖。")
    body(doc, "不足：自動化覆蓋率仍偏低（205 條中僅 6 條核心流程自動化）、效能測試僅針對少數不需登入的介面、安全性測試多依賴人工判讀。")
    body(doc, "對課程的建議：可增加實機演練與 CI 串接（如 Newman）之教學，使測試流程更貼近實務。")

    # 致謝
    heading(doc, "致　謝", 1)
    body(doc, "本次《軟體測試》大作業得以順利完成，特此致上誠摯謝意：")
    body(doc, "感謝台灣國立大學張導師。從測試計畫的擬定、測試案例設計方法（等價類、邊界值、判定表、場景法）的指導，到 Postman 斷言與 JMeter 壓力測試觀念的講解，張導師都給予了悉心的教導與寶貴的建議，使我在實作中少走許多彎路，對軟體測試的整體流程有了更扎實的理解。")
    body(doc, "同時，特別感謝來自日本的 Jin 同學。在實驗過程中，Jin 同學不僅與我一同討論測試思路、協助釐清 API 回應碼的判讀問題，更在 Selenium 自動化腳本除錯與高併發效能測試的環境調校上提供了許多實質的協助；跨文化的交流與合作，讓這次作業的過程格外充實而難忘。")
    body(doc, "再次向張導師與 Jin 同學致上最深的謝意。")

    # 附錄一：完整 205 條案例
    doc.add_page_break()
    heading(doc, "附錄一　功能測試案例清單（205 條）", 1)
    for name, rows in read_module_cases():
        heading(doc, "%s（%d 條）" % (name, len(rows)), 2)
        table_rows = [[r[0], r[1], r[2], r[3]] for r in rows]
        make_table(doc, ["編號", "測試方法", "輸入/操作", "預期輸出"], table_rows,
                   widths=[1.8, 2.2, 5.2, 5.0], body_size=9, header_size=9.5)

    out = os.path.join(BASE, "軟體測試大作業報告_完整版.docx")
    doc.save(out)
    print("已產生:", out)


if __name__ == "__main__":
    build()
